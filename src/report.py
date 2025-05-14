import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Side
import matplotlib.pyplot as plt
import gspread
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
import os
from dotenv import load_dotenv
from datetime import datetime
from gspread_formatting import format_cell_range, CellFormat, Color, TextFormat, Border as GSpreadBorder, Borders

load_dotenv()

# Google Sheets configuration
SCOPES = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
CREDS_FILE = os.getenv("GOOGLE_SHEETS_CREDENTIALS")


def create_excel_report(weather_data, output_file, style="minimal"):
    """Create a formatted Excel report with specified style."""
    from openpyxl.styles import Border  # Explicit import to avoid conflict

    if not weather_data:
        print("No weather data provided. Skipping Excel report creation.")
        return

    df = pd.DataFrame(weather_data)

    # Create workbook and worksheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weather Report"

    # Style configurations
    styles = {
        "minimal": {
            "title_bg": "003087", "title_fg": "FFFFFF",
            "header_bg": "D3D3D3", "header_fg": "000000",
            "font": "Arial", "title_size": 18, "header_size": 12,
            "chart_type": "line"
        },
        "modern": {
            "title_bg": "28A745", "title_fg": "FFFFFF",
            "header_bg": "007BFF", "header_fg": "FFFFFF",
            "font": "Roboto", "title_size": 20, "header_size": 14,
            "chart_type": "bar"
        },
        "corporate": {
            "title_bg": "333333", "title_fg": "FFFFFF",
            "header_bg": "4BACC6", "header_fg": "FFFFFF",
            "font": "Times New Roman", "title_size": 16, "header_size": 11,
            "chart_type": "combo"
        }
    }
    s = styles.get(style, styles["minimal"])

    # Add title
    ws.merge_cells("A1:H1")
    ws["A1"] = f"Weather Report - {datetime.now().strftime('%Y-%m-%d')}"
    ws["A1"].font = Font(name=s["font"], size=s["title_size"], bold=True, color=s["title_fg"])
    ws["A1"].fill = PatternFill(start_color=s["title_bg"], end_color=s["title_bg"], fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center")

    # Add headers
    headers = ["City", "Date", "Time", "Temperature (°C)", "Precipitation (mm)",
               "Humidity (%)", "Pressure (hPa)", "Wind Speed (m/s)"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col)
        cell.value = header
        cell.font = Font(name=s["font"], size=s["header_size"], bold=True, color=s["header_fg"])
        cell.fill = PatternFill(start_color=s["header_bg"], end_color=s["header_bg"], fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        border = Border()
        border.left = Side(border_style="thin")
        border.right = Side(border_style="thin")
        border.top = Side(border_style="thin")
        border.bottom = Side(border_style="thin")
        cell.border = border

    # Add data with city headers
    row = 4
    current_city = None
    for data in weather_data:
        if data["city"] != current_city:
            current_city = data["city"]
            ws.cell(row=row, column=1).value = f"{current_city} Weather"
            ws.merge_cells(f"A{row}:H{row}")
            ws.cell(row=row, column=1).font = Font(name=s["font"], size=s["header_size"], bold=True)
            ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
            row += 1

        ws.cell(row=row, column=1).value = data["city"]
        ws.cell(row=row, column=2).value = data["date"]
        ws.cell(row=row, column=3).value = data["time"]
        ws.cell(row=row, column=4).value = data["temperature"]
        ws.cell(row=row, column=5).value = data["precipitation"]
        ws.cell(row=row, column=6).value = data["humidity"]
        ws.cell(row=row, column=7).value = data["pressure"]
        ws.cell(row=row, column=8).value = data["wind_speed"]

        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            border = Border()
            border.left = Side(border_style="thin")
            border.right = Side(border_style="thin")
            border.top = Side(border_style="thin")
            border.bottom = Side(border_style="thin")
            cell.border = border
        row += 1

    # Create chart
    if len(df) > 1:  # Ensure enough data for chart
        plt.figure(figsize=(12, 6))  # Increased width to accommodate legend

        # Colors for each city
        city_colors = {
            "Moscow": "#FF0000",  # Red
            "Chisinau": "#0000FF",  # Blue
            "Dublin": "#00FF00",  # Green
            "London": "#FFA500",  # Orange
            "Berlin": "#800080",  # Purple
            "Paris": "#00FFFF"  # Cyan
        }

        # Group data by city
        cities = df["city"].unique()

        if s["chart_type"] == "line":
            for city in cities:
                city_data = df[df["city"] == city]
                plt.plot(city_data["date"], city_data["temperature"],
                         color=city_colors.get(city, "#000000"), marker="o",
                         label=f"{city} Temperature")
            plt.legend(bbox_to_anchor=(1.05, 1), loc='upper left', borderaxespad=0.)

        elif s["chart_type"] == "bar":
            bar_width = 0.15
            x = range(len(df["date"].unique()))
            for i, city in enumerate(cities):
                city_data = df[df["city"] == city]
                plt.bar([xi + i * bar_width for xi in x], city_data["temperature"],
                        width=bar_width, color=city_colors.get(city, "#000000"),
                        label=f"{city} Temperature")
            plt.xticks([xi + bar_width * (len(cities) - 1) / 2 for xi in x], df["date"].unique(), rotation=45)
            plt.legend(bbox_to_anchor=(1.05, 1), loc='upper left', borderaxespad=0.)

        else:  # combo
            fig, ax1 = plt.subplots(figsize=(12, 6))
            ax2 = ax1.twinx()  # Secondary axis for precipitation
            temp_handles = []
            precip_handles = []
            for city in cities:
                city_data = df[df["city"] == city]
                # Temperature (line)
                line, = ax1.plot(city_data["date"], city_data["temperature"],
                                 color=city_colors.get(city, "#000000"), marker="o",
                                 label=f"{city} Temperature")
                temp_handles.append(line)
                # Precipitation (bar)
                bar = ax2.bar(city_data["date"], city_data["precipitation"],
                              color=city_colors.get(city, "#000000"), alpha=0.3,
                              label=f"{city} Precipitation", width=0.15)
                precip_handles.append(bar)

            # Create two legends
            legend1 = ax1.legend(handles=temp_handles, bbox_to_anchor=(1.05, 1), loc='upper left', title="Temperature")
            ax1.add_artist(legend1)
            legend2 = ax2.legend(handles=precip_handles, bbox_to_anchor=(1.05, 0.5), loc='center left',
                                 title="Precipitation")
            ax1.set_xlabel("Date")
            ax1.set_ylabel("Temperature (°C)")
            ax2.set_ylabel("Precipitation (mm)")
            plt.title("Weather Forecast by City")
            ax1.grid(True, linestyle="--", alpha=0.7)
            plt.xticks(rotation=45)
            plt.tight_layout()

        plt.savefig("weather_chart.png", bbox_inches='tight')
        plt.close()

        # Add chart to Excel
        img = openpyxl.drawing.image.Image("weather_chart.png")
        ws.add_image(img, "A" + str(row + 2))
    else:
        print("Not enough data to create chart")

    # Adjust column widths
    column_widths = {"A": 15, "B": 15, "C": 10, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Save the workbook
    wb.save(output_file)
    print(f"Excel report saved as {output_file} (style: {style})")


def create_gsheets_report(weather_data, spreadsheet_name="Weather Reports"):
    """Create a formatted Google Sheets report with corporate style and automated chart."""
    if not weather_data:
        print("No weather data provided. Skipping Google Sheets report creation.")
        return

    # Initialize gspread client
    creds = Credentials.from_service_account_file(CREDS_FILE, scopes=SCOPES)
    client = gspread.authorize(creds)

    try:
        spreadsheet = client.open(spreadsheet_name)
    except gspread.exceptions.SpreadsheetNotFound:
        spreadsheet = client.create(spreadsheet_name)
        spreadsheet.share(os.getenv("GOOGLE_SHEETS_CLIENT_EMAIL"), perm_type="user", role="writer")

    worksheet = spreadsheet.sheet1
    worksheet.clear()

    # Add title
    title = f"Weather Report - {datetime.now().strftime('%Y-%m-%d')}"
    worksheet.update("A1:H1", [[title]])
    format_cell_range(worksheet, "A1:H1", CellFormat(
        textFormat=TextFormat(fontSize=16, bold=True, foregroundColor=Color(1, 1, 1)),
        backgroundColor=Color(0.2, 0.2, 0.2),
        horizontalAlignment="CENTER"
    ))
    worksheet.merge_cells("A1:H1")

    # Add headers
    headers = ["City", "Date", "Time", "Temperature (°C)", "Precipitation (mm)",
               "Humidity (%)", "Pressure (hPa)", "Wind Speed (m/s)"]
    worksheet.update("A3:H3", [headers])
    format_cell_range(worksheet, "A3:H3", CellFormat(
        textFormat=TextFormat(bold=True, foregroundColor=Color(1, 1, 1)),
        backgroundColor=Color(0.29, 0.67, 0.78),
        horizontalAlignment="CENTER",
        borders=Borders(
            top=GSpreadBorder("SOLID"),
            bottom=GSpreadBorder("SOLID"),
            left=GSpreadBorder("SOLID"),
            right=GSpreadBorder("SOLID")
        )
    ))

    # Add data with city headers
    row = 4
    current_city = None
    data_rows = []
    for data in weather_data:
        if data["city"] != current_city:
            current_city = data["city"]
            data_rows.append([f"{current_city} Weather"] + [""] * 7)
            row += 1
        data_rows.append([
            data["city"], data["date"], data["time"], data["temperature"],
            data["precipitation"], data["humidity"], data["pressure"], data["wind_speed"]
        ])
        row += 1

    worksheet.update(f"A4:H{row - 1}", data_rows)
    format_cell_range(worksheet, f"A4:H{row - 1}", CellFormat(
        borders=Borders(
            top=GSpreadBorder("SOLID"),
            bottom=GSpreadBorder("SOLID"),
            left=GSpreadBorder("SOLID"),
            right=GSpreadBorder("SOLID")
        )
    ))

    # Prepare chart data in a pivoted format
    df = pd.DataFrame(weather_data)
    cities = ["Moscow", "Chisinau", "Dublin", "London", "Berlin", "Paris"]
    dates = sorted(df["date"].unique())
    chart_data = [["Date"] + [f"{city} Temp" for city in cities] + [f"{city} Precip" for city in cities]]

    for date in dates:
        row = [date]
        for city in cities:
            city_data = df[(df["city"] == city) & (df["date"] == date)]
            temp = city_data["temperature"].iloc[0] if not city_data.empty else 0
            row.append(temp)
        for city in cities:
            city_data = df[(df["city"] == city) & (df["date"] == date)]
            precip = city_data["precipitation"].iloc[0] if not city_data.empty else 0
            row.append(precip)
        chart_data.append(row)

    print(f"Chart data: {chart_data}")  # Debug output
    worksheet.update(f"I4:U{4 + len(dates)}", chart_data)

    # Initialize Google Sheets API client for chart creation
    sheets_service = build('sheets', 'v4', credentials=creds)
    spreadsheet_id = spreadsheet.id

    # Colors for each city (RGB values normalized to 0-1)
    city_colors = {
        "Moscow": {"red": 1.0, "green": 0.0, "blue": 0.0},  # Red
        "Chisinau": {"red": 0.0, "green": 0.0, "blue": 1.0},  # Blue
        "Dublin": {"red": 0.0, "green": 1.0, "blue": 0.0},  # Green
        "London": {"red": 1.0, "green": 0.647, "blue": 0.0},  # Orange
        "Berlin": {"red": 0.502, "green": 0.0, "blue": 0.502},  # Purple
        "Paris": {"red": 0.0, "green": 1.0, "blue": 1.0}  # Cyan
    }

    # Create chart request
    requests = []
    series = []

    # Temperature series (lines)
    for i, city in enumerate(cities):
        series.append({
            "series": {
                "sourceRange": {
                    "sources": [{
                        "sheetId": worksheet.id,
                        "startRowIndex": 3,
                        "endRowIndex": 3 + len(dates),
                        "startColumnIndex": 9 + i,  # J:O (Temperature columns)
                        "endColumnIndex": 10 + i
                    }]
                }
            },
            "colorStyle": {
                "rgbColor": city_colors.get(city, {"red": 0, "green": 0, "blue": 0})
            },
            "targetAxis": "LEFT_AXIS",
            "type": "LINE"
        })

    # Precipitation series (columns)
    for i, city in enumerate(cities):
        series.append({
            "series": {
                "sourceRange": {
                    "sources": [{
                        "sheetId": worksheet.id,
                        "startRowIndex": 3,
                        "endRowIndex": 3 + len(dates),
                        "startColumnIndex": 15 + i,  # P:U (Precipitation columns)
                        "endColumnIndex": 16 + i
                    }]
                }
            },
            "colorStyle": {
                "rgbColor": city_colors.get(city, {"red": 0, "green": 0, "blue": 0})
            },
            "targetAxis": "RIGHT_AXIS",
            "type": "COLUMN"
        })

    # Chart request
    requests.append({
        "addChart": {
            "chart": {
                "spec": {
                    "title": "Weather Forecast by City",
                    "basicChart": {
                        "chartType": "COMBO",
                        "legendPosition": "RIGHT_LEGEND",
                        "axis": [
                            {
                                "position": "BOTTOM_AXIS",
                                "title": "Date"
                            },
                            {
                                "position": "LEFT_AXIS",
                                "title": "Temperature (°C)"
                            },
                            {
                                "position": "RIGHT_AXIS",
                                "title": "Precipitation (mm)"
                            }
                        ],
                        "domains": [{
                            "domain": {
                                "sourceRange": {
                                    "sources": [{
                                        "sheetId": worksheet.id,
                                        "startRowIndex": 3,
                                        "endRowIndex": 3 + len(dates),
                                        "startColumnIndex": 8,  # I (Date)
                                        "endColumnIndex": 9
                                    }]
                                }
                            }
                        }],
                        "series": series,
                        "headerCount": 1  # Use first row (I4:U4) for series labels
                    }
                },
                "position": {
                    "overlayPosition": {
                        "anchorCell": {
                            "sheetId": worksheet.id,
                            "rowIndex": 50,
                            "columnIndex": 0
                        },
                        "offsetXPixels": 10,
                        "offsetYPixels": 10,
                        "widthPixels": 600,
                        "heightPixels": 400
                    }
                }
            }
        }
    })

    # Execute chart creation
    try:
        print(f"Chart requests: {requests}")  # Debug output
        response = sheets_service.spreadsheets().batchUpdate(
            spreadsheetId=spreadsheet_id,
            body={"requests": requests}
        ).execute()
        print(f"Chart creation response: {response}")  # Debug output
    except HttpError as e:
        print(f"Error creating chart: {e}")
        raise

    print(f"Google Sheets report updated: {spreadsheet_name} with chart")